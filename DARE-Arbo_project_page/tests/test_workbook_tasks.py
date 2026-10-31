from io import BytesIO
from hashlib import sha256
import time
import unittest
from unittest.mock import MagicMock, patch
from openpyxl import Workbook, load_workbook
from workbook_tasks import WorkbookTasks
from send_task_reminders import eligible
from datetime import date


class WorkbookTaskTests(unittest.TestCase):
    def setUp(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Activity Register'
        sheet.append(['ID', 'Workstream', 'Activity', 'Key output / milestone', 'Start', 'End', 'Owner', 'Status', 'Assignee email', 'Progress (%)'])
        sheet.append([1, 'Pilot', 'Review', 'Dataset', '2026-09-01', '2026-10-31', 'EAA', 'In progress', 'one@example.org', 25])
        other = book.create_sheet('Gantt')
        other['A1'] = '=1+1'
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        self.content = buffer.getvalue()
        self.person = dict(email='admin@example.org', is_logged_in=True, email_verified=True,
                           iss='https://accounts.google.com', exp=time.time()+3600)
        self.people = ['admin@example.org', 'one@example.org', 'two@example.org']
        self.service = MagicMock()
        self.store = WorkbookTasks(self.service, lambda: self.person, lambda: self.people, lambda: self.people[:1])
        self.store.version_service = self.service
        self.store.read = MagicMock(side_effect=lambda: (self.content, 'etag'))
        def update(**kwargs):
            request = MagicMock()
            request.headers = {}
            def execute():
                media = kwargs['media_body']
                self.assertEqual('etag', request.headers['If-Match'])
                self.content = media.getbytes(0, media.size())
                return {'id': 'source'}
            request.execute.side_effect = execute
            return request
        self.service.files.return_value.update.side_effect = update

    def test_collaborator_can_edit_only_own_progress(self):
        self.person['email'] = self.people[2]
        self.assertEqual([], self.store.list())
        with self.assertRaises(PermissionError):
            self.store.save({'status': 'Completed'}, 2, sha256(self.content).hexdigest())
        self.person['email'] = self.people[1]
        task = self.store.list()[0]
        with self.assertRaises(PermissionError):
            self.store.save({'assignee': self.people[2]}, 2, task['version'])
        self.store.save({'status': 'Completed', 'progress': 50, 'note': '=literal'}, 2, task['version'])
        self.assertEqual(100, self.store.list()[0]['progress'])
        book = load_workbook(BytesIO(self.content))
        self.assertEqual('=1+1', book['Gantt']['A1'].value)
        note = next(c.column for c in book['Activity Register'][1] if c.value == 'Progress notes')
        self.assertEqual('s', book['Activity Register'].cell(2, note).data_type)
        book.close()

    def test_admin_creation_and_stale_version(self):
        fields = dict(title='New task', description='Output', owner='AO', workstream='Pilot',
                      assignee=self.people[2], start_date='2026-10-01', due_date='2026-11-30',
                      status='Not started', progress=0, note='')
        stale = self.store.list()[0]['version']
        self.store.save(fields)
        self.assertEqual(2, len(self.store.list()))
        with self.assertRaises(ValueError):
            self.store.save({'progress': 20}, 2, stale)
        with self.assertRaises(ValueError):
            self.store.save({**fields, 'assignee': 'unknown@example.org'})
        self.person['email_verified'] = False
        with self.assertRaises(PermissionError):
            self.store.list()

    def test_reminder_uses_workbook_and_claims_once(self):
        task = self.store.list()[0]
        self.assertTrue(eligible(task, self.people, date(2026, 10, 30)))
        self.assertFalse(eligible(task, [], date(2026, 10, 30)))
        self.assertTrue(self.store.claim_reminder(2, '2026-10-30'))
        self.assertFalse(self.store.claim_reminder(2, '2026-10-30'))

    def test_no_write_without_etag(self):
        self.store.read.return_value = None
        self.store.read.side_effect = lambda: (self.content, None)
        with self.assertRaises(ValueError):
            self.store.save({'progress': 50}, 2, sha256(self.content).hexdigest())
        self.service.files.return_value.update.assert_not_called()

    def test_conditional_write_conflict_is_reported(self):
        from googleapiclient.errors import HttpError
        from httplib2 import Response
        task = self.store.list()[0]
        request = MagicMock()
        request.headers = {}
        request.execute.side_effect = HttpError(Response({'status':'412'}), b'conflict')
        self.service.files.return_value.update.side_effect = None
        self.service.files.return_value.update.return_value = request
        with self.assertRaisesRegex(ValueError, 'workbook changed'):
            self.store.save({'progress':40},2,task['version'])

    def test_progress_update_preserves_invalid_dates_and_unrelated_fields(self):
        book = load_workbook(BytesIO(self.content))
        book['Activity Register']['E2'] = 'Not yet specified'
        book['Activity Register']['J2'] = None
        out = BytesIO(); book.save(out); book.close()
        self.content = out.getvalue()
        self.person['email'] = self.people[1]
        original = self.store.list()[0]
        self.assertIsNone(original['progress'])
        self.store.save({'note': 'Awaiting date confirmation'}, 2, original['version'])
        book = load_workbook(BytesIO(self.content))
        self.assertEqual('Not yet specified', book['Activity Register']['E2'].value)
        self.assertIsNone(book['Activity Register']['J2'].value)
        self.assertEqual('=1+1', book['Gantt']['A1'].value)
        book.close()

    def test_project_page_uses_only_workbook(self):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_string('''
from project_management import show_project_management
import streamlit as st
show_project_management(lambda: {'drive': dict(client_id='test', client_secret='test', refresh_token='test'), 'project_management': {'admin_emails': ['admin@example.org']}},
 lambda: st.session_state['person'], lambda: ['admin@example.org','one@example.org'])
''', default_timeout=30)
        app.session_state['person'] = self.person.copy()
        with patch('project_management.create_service'), patch('sheet_dashboard.create_service'), patch('sheet_dashboard.DriveStore.download', side_effect=lambda *args: ('file.xlsx','',self.content)) as download, patch('project_management.WorkbookTasks', return_value=self.store):
            app.run()
            self.assertFalse(app.exception)
            self.assertFalse(app.error)
            self.assertTrue(any(t.label == 'Assignee email' for t in app.text_input))
            self.assertEqual(['Overview', 'Activity Register', 'Study Timeline', 'Task Management'], [t.label for t in app.tabs])
            self.assertEqual(1, download.call_count)
            app.session_state['person'] = {**self.person, 'email': self.people[1]}
            self.person['email'] = self.people[1]
            app.run()
            self.assertFalse(app.exception)
            self.assertFalse(any(t.label == 'Assignee email' for t in app.text_input))
            self.assertEqual('My Tasks', app.tabs[-1].label)
            next(b for b in app.button if b.label == 'Open activity editor').click().run()
            self.assertFalse(app.exception)
            next(n for n in app.number_input if n.label == 'Recorded progress (%)').set_value(50.0)
            next(b for b in app.button if b.label == 'Keep draft and choose another').click().run()
            self.assertFalse(app.exception)
            next(b for b in app.button if b.label == 'Open activity editor').click().run()
            self.assertEqual(50, next(n for n in app.number_input if n.label == 'Recorded progress (%)').value)
            next(b for b in app.button if b.label == 'Save update').click().run()
            self.assertFalse(app.exception)
            self.assertFalse(app.error)
            self.assertEqual(50, self.store.list()[0]['progress'])
            # A draft retains its displayed version despite another source edit.
            next(b for b in app.button if b.label == 'Open activity editor').click().run()
            book=load_workbook(BytesIO(self.content));book['Gantt']['B1']='Another editor'
            out=BytesIO();book.save(out);book.close();self.content=out.getvalue()
            next(t for t in app.text_area if t.label == 'Progress notes').set_value('New notes')
            next(b for b in app.button if b.label == 'Save update').click().run()
            self.assertFalse(app.exception)
            self.assertTrue(any('workbook changed' in e.value for e in app.error))

    def test_admin_creates_activity_through_form(self):
        from streamlit.testing.v1 import AppTest
        app=AppTest.from_string('''
from project_management import show_project_management
import streamlit as st
show_project_management(lambda: {'drive':dict(client_id='test',client_secret='test',refresh_token='test'),
 'project_management':{'admin_emails':['admin@example.org']}}, lambda:st.session_state['person'],
 lambda:['admin@example.org','one@example.org'])
''',default_timeout=30)
        app.session_state['person']=self.person.copy()
        with patch('project_management.create_service'), patch('sheet_dashboard.create_service'), patch('sheet_dashboard.DriveStore.download',side_effect=lambda *a:('file.xlsx','',self.content)), patch('project_management.WorkbookTasks',return_value=self.store):
            app.run()
            for label,value in [('Activity title','New pilot review'),('Workstream','Pilot'),('Responsible owner','EAA'),('Assignee email','one@example.org')]:
                next(t for t in app.text_input if t.label==label).set_value(value)
            next(t for t in app.date_input if t.label=='Start date').set_value(date(2026,11,1))
            next(t for t in app.date_input if t.label=='Deadline').set_value(date(2026,11,30))
            next(b for b in app.button if b.label=='Save activity').click().run()
            self.assertFalse(app.exception)
            self.assertFalse(app.error)
            tasks=self.store.list()
            self.assertEqual(2,len(tasks))
            self.assertEqual('New pilot review',tasks[-1]['title'])
            self.assertIsNone(tasks[-1]['progress'])


if __name__ == '__main__':
    unittest.main()
