from datetime import datetime,timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from unittest import TestCase
from unittest.mock import MagicMock,patch
from PIL import Image
from openpyxl import load_workbook
import test_workbook_tasks as tasks
from team_store import TeamProfiles,parse_team,normalize_portrait,public_blob,MIME
from drive_store import FOLDER_ID
from assignment_notifications import verify_sender


class TeamProfileTests(TestCase):
    def setUp(self):
        tasks.WorkbookTaskTests.setUp(self)
        book=load_workbook(BytesIO(self.content))
        sheet=book.create_sheet('Team_members')
        sheet.append(['Name','Initial','Affiliation','Experience','Position','Private email'])
        sheet.append(['Member one','MO','Institution','Surveillance','Researcher','private@example.org'])
        buffer=BytesIO();book.save(buffer);book.close();self.content=buffer.getvalue()
        self.team=TeamProfiles(self.service,lambda:self.person,lambda:self.people,lambda:self.people[:1])
        self.team.version_service=self.service
        self.team.read=MagicMock(side_effect=lambda:(self.content,'etag'))
        self.fields=dict(name='Member one',initials='MO',affiliation='Institution',expertise='Surveillance',position='Researcher',featured=True)

    def test_admin_can_add_edit_and_preserve_workbook(self):
        self.team.save_profile({**self.fields,'name':'New member'},version=sha256(self.content).hexdigest())
        self.assertEqual(2,len(parse_team(self.content)))
        self.team.save_profile({**self.fields,'expertise':'=literal text'},row=2,version=sha256(self.content).hexdigest())
        book=load_workbook(BytesIO(self.content))
        self.assertEqual('=1+1',book['Gantt']['A1'].value)
        self.assertEqual('s',book['Team_members']['D2'].data_type)
        self.assertEqual('private@example.org',book['Team_members']['F2'].value)
        book.close()
        self.assertNotIn('private@example.org',str(parse_team(self.content)))

    def test_nonadmin_and_stale_edit_rejected(self):
        version=sha256(self.content).hexdigest()
        self.person['email']=self.people[1]
        with self.assertRaises(PermissionError):self.team.save_profile(self.fields,2,version)
        self.person['email']=self.people[0]
        self.team.save_profile(self.fields,2,version)
        with self.assertRaises(ValueError):self.team.save_profile(self.fields,2,version)

    def test_image_validation_and_metadata_restrictions(self):
        image=Image.new('RGB',(1000,600),'white');out=BytesIO();image.save(out,format='JPEG')
        normalized=normalize_portrait(out.getvalue())
        with Image.open(BytesIO(normalized)) as result:
            self.assertEqual('PNG',result.format);self.assertLessEqual(result.width,800)
        with self.assertRaises(ValueError):normalize_portrait(b'<svg/>')
        with self.assertRaises(PermissionError):public_blob(self.service,'not-the-team-source')
        self.service.files.return_value.get.return_value.execute.return_value=dict(parents=[FOLDER_ID],mimeType='image/png',size=100,appProperties={})
        with self.assertRaises(PermissionError):public_blob(self.service,'private-image',portrait=True)

    def test_photo_is_referenced_only_after_admin_save(self):
        image=Image.new('RGB',(40,60),'white');out=BytesIO();image.save(out,format='PNG')
        self.service.files.return_value.create.return_value.execute.return_value={'id':'published-photo'}
        self.team.save_profile(self.fields,2,sha256(self.content).hexdigest(),out.getvalue())
        self.assertEqual('published-photo',parse_team(self.content)[0]['photo_id'])
        body=self.service.files.return_value.create.call_args.kwargs['body']
        self.assertEqual('public_team_portrait',body['appProperties']['purpose'])

    def test_public_refresh_and_admin_editor(self):
        from streamlit.testing.v1 import AppTest
        script='''
from pathlib import Path
import streamlit as st
from team_profiles import show_team
show_team(lambda:{'drive':dict(client_id='test',client_secret='test',refresh_token='test'),
'project_management':{'admin_emails':['admin@example.org']}},lambda:st.session_state['person'],
lambda:['admin@example.org','one@example.org'],Path.cwd(),{'team':[]})
'''
        app=AppTest.from_string(script,default_timeout=30)
        app.session_state['person']={}
        def fetch(config):
            return dict(members=parse_team(self.content),version=sha256(self.content).hexdigest(),synced=datetime.now(timezone.utc))
        with patch('team_profiles.fetch_public_team',side_effect=fetch),patch('team_profiles.create_service'),patch('team_profiles.TeamProfiles',return_value=self.team):
            app.run()
            self.assertFalse(app.exception)
            self.assertFalse(app.file_uploader)
            self.team.save_profile({**self.fields,'name':'Added member'},version=sha256(self.content).hexdigest())
            app.session_state['team_snapshot']['loaded']=0
            app.run()
            self.assertTrue(any('2 team members' in c.value for c in app.caption))
            app.session_state['person']=self.person.copy()
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(1,len(app.file_uploader))
            next(t for t in app.text_input if t.label=='Full name').set_value('UI member')
            next(t for t in app.text_input if t.label=='Position / role').set_value('Researcher')
            next(t for t in app.text_input if t.label=='Institution / affiliation').set_value('Institution')
            next(b for b in app.button if b.label=='Publish team profile').click().run()
            self.assertFalse(app.exception)
            self.assertFalse(app.error)
            self.assertEqual(3,len(parse_team(self.content)))

    def test_smtp_verification_does_not_send(self):
        config=dict(host='smtp.test',port=587,username='admin@example.org',password='test',from_address='admin@example.org',app_url='https://example.org')
        with patch('assignment_notifications.smtplib.SMTP') as smtp:
            verify_sender(config)
            smtp.return_value.__enter__.return_value.login.assert_called_once()
            smtp.return_value.__enter__.return_value.send_message.assert_not_called()
