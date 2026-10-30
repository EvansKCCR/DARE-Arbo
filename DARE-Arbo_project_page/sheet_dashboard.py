"""Live, read-only study dashboard backed by the project Drive workbook."""
from datetime import date, datetime, timezone
from html import escape
from io import BytesIO
import logging
import time
from hashlib import sha256
import streamlit as st
from openpyxl import load_workbook
from google.auth.exceptions import RefreshError
from access import has_full_access
from drive_store import DriveStore, create_service

SOURCE_ID = '1Pmf6ZYq3HyooclARiVtZFqYtI6f-sNQx'
SOURCE_URL = f'https://docs.google.com/spreadsheets/d/{SOURCE_ID}/edit#gid=1783506628'


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or '').strip().replace('Sept', 'Sep')
    for pattern in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d-%b-%y', '%d %b %Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError('Unrecognized date')


def parse_register(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook['Activity Register'] if 'Activity Register' in workbook.sheetnames else None
        if sheet is None:
            raise ValueError('The workbook needs an Activity Register tab.')
        header = None
        tasks, issues = [], []
        for number, cells in enumerate(sheet.iter_rows(), 1):
            row = [cell.value for cell in cells]
            labels = [str(v or '').strip().casefold() for v in row]
            if header is None:
                if {'activity', 'start', 'end', 'owner', 'status'}.issubset(labels):
                    header = {label: index for index, label in enumerate(labels) if label}
                continue
            def get(key, default=''):
                index = header.get(key)
                return row[index] if index is not None and index < len(row) and row[index] is not None else default
            if not str(get('activity')).strip():
                continue
            dates = []
            for name in ('start', 'end'):
                try:
                    dates.append(parse_date(get(name)))
                except ValueError:
                    dates.append(None)
            start, end = dates
            if not start or not end or end < start:
                issues.append(f'Row {number}: missing or invalid start/end dates; check the recorded schedule.')
            status = str(get('status', 'Not recorded')).strip()
            raw = get('progress (%)', get('progress', get('% complete', '')))
            progress = None
            if str(raw).strip():
                try:
                    progress = float(str(raw).replace('%', '').strip())
                    progress_index = next((header[k] for k in ('progress (%)', 'progress', '% complete') if k in header), None)
                    if isinstance(raw, (float, int)) and progress_index is not None and '%' in cells[progress_index].number_format:
                        progress *= 100
                    if not 0 <= progress <= 100:
                        raise ValueError()
                except ValueError:
                    issues.append(f'Row {number}: invalid progress percentage; using status only.')
                    progress = None
            tasks.append(dict(id=str(get('id')).strip() or 'Not recorded', activity=str(get('activity')).strip(),
                              row=number, assignee=str(get('assignee email')).strip(), note=str(get('progress notes')),
                              reminder_day=str(get('last reminder date')), reminder_state=str(get('last reminder state')),
                              workstream=str(get('workstream', 'Other')).strip(),
                              output=str(get('key output / milestone')), owner=str(get('owner', 'Unassigned')),
                              status=status, start=start, end=end, progress=progress))
        if header is None:
            raise ValueError('Activity Register headers were not found.')
        from collections import Counter
        counts = Counter(t['id'] for t in tasks)
        for task in tasks:
            valid_id = task['id'] != 'Not recorded' and counts[task['id']] == 1
            task['record_key'] = task['id'] if valid_id else f'row:{task["row"]}'
            if not valid_id:
                issues.append(f'Row {task["row"]}: missing or duplicate activity ID; assign a unique ID for stable selection.')
        return tasks, issues
    finally:
        workbook.close()


def elapsed(start, end, today):
    return max(0.0, min(1.0, (today - start).days / max(1, (end - start).days)))



def snapshot_scope(settings, identity, allowed):
    from access import require_full_access
    require_full_access(identity(), allowed())
    cfg = settings()
    source = str(sorted(cfg.get('drive', {}).items()))
    roles = str(cfg.get('project_management', {}).get('admin_emails', ['evansasamoahadu@gmail.com']))
    return sha256((identity()['email'].strip().casefold() + str(sorted(allowed())) + source + roles).encode()).hexdigest()


def load_snapshot(settings, identity, allowed, refresh=False, allow_expired=False):
    scope = snapshot_scope(settings, identity, allowed)
    if st.session_state.get('pm_scope') != scope:
        for key in list(st.session_state):
            if key.startswith(('pm_', '_pm_', 'workbook_task_')):
                del st.session_state[key]
        st.session_state['pm_scope'] = scope
    cached = st.session_state.get('pm_snapshot')
    if cached and not refresh and (allow_expired or time.monotonic() - cached['loaded'] < 60):
        return cached
    config = settings().get('drive', {})
    if not all(config.get(k) for k in ('client_id', 'client_secret', 'refresh_token')):
        raise ValueError('Drive connection is not configured.')
    try:
        _, _, content = DriveStore(create_service(config), identity, allowed).download(SOURCE_ID)
        tasks, issues = parse_register(content)
        from assignment_notifications import read_events
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            notifications = read_events(book)
        finally:
            book.close()
    except Exception:
        st.session_state.pop('pm_snapshot', None)
        raise
    result = dict(tasks=tasks, issues=issues, notifications=notifications, version=sha256(content).hexdigest(),
                  loaded=time.monotonic(), synced=datetime.now(timezone.utc))
    st.session_state['pm_snapshot'] = result
    return result


def invalidate_snapshot():
    st.session_state.pop('pm_snapshot', None)


def sync_header(snapshot, settings, identity, allowed, key):
    left, middle, right = st.columns([3, 2, 1])
    left.caption('Workbook synchronized · ' + snapshot['synced'].strftime('%d %b %Y %H:%M:%S UTC'))
    middle.link_button('Open source workbook', SOURCE_URL, width='stretch')
    right.button('Refresh', key='pm_refresh_' + key, width='stretch', on_click=invalidate_snapshot)
    if snapshot['issues']:
        with st.expander(f"Workbook data checks · {len(snapshot['issues'])} items"):
            for issue in snapshot['issues']:
                st.caption(issue)


@st.fragment(run_every='60s')
def show_sheet_dashboard(settings, identity, allowed, view='Overview'):
    if not has_full_access(identity(), allowed()):
        st.info('Sign in with an approved collaborator account to view the study workspace.')
        return
    try:
        snapshot = load_snapshot(settings, identity, allowed)
    except RefreshError:
        st.error('Google Drive authorization has expired or been revoked. Ask the administrator to reconnect the folder-owner account.')
        return
    except Exception:
        logging.getLogger(__name__).warning('Study workbook could not be synchronized.')
        st.error('The workbook could not be synchronized. Check the Drive connection and Activity Register headers, then retry.')
        st.button('Retry synchronization', key='pm_retry_' + view)
        return
    sync_header(snapshot, settings, identity, allowed, view)
    from workspace_views import overview, register, study_timeline
    # Each read section owns its timed fragment; task forms are outside them.
    with st.container(key='pm_read_' + view.replace(' ', '_')):
        render = {'Overview': overview, 'Activity Register': register, 'Study Timeline': study_timeline}[view]
        if view == 'Study Timeline':
            render(snapshot['tasks'], datetime.now(timezone.utc).date(), identity()['email'])
        else:
            render(snapshot['tasks'], datetime.now(timezone.utc).date())
    st.caption('Automatic synchronization every 60 seconds while this workspace is open. Task drafts are not refreshed automatically.')
