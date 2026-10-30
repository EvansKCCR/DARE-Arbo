"""Authorized task edits in the single Drive-hosted project workbook."""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import re
from openpyxl import load_workbook
from googleapiclient.http import MediaIoBaseUpload
from access import has_full_access
from drive_store import FOLDER_ID, MAX_BYTES
from sheet_dashboard import SOURCE_ID, parse_register, parse_date
from assignment_notifications import initialize_tracking, record_assignment, read_events, TRACKING, QUEUE

STATUSES = ['Not started', 'Planned / update', 'In progress', 'Under review', 'Completed', 'Blocked']
PRIORITIES = ['Low', 'Normal', 'High']
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def emails(value):
    return {s.strip().casefold() for s in re.split('[,;\\n]', value) if s.strip()}


class WorkbookTasks:
    def __init__(self, service, identity, allowed, admins):
        self.service, self.identity, self.allowed, self.admins = service, identity, allowed, admins
        self.version_service = None

    def versioned(self):
        if self.version_service is None:
            from googleapiclient.discovery import build
            # v2 exposes file ETags explicitly; v3 omits them on this workbook.
            self.version_service = build('drive', 'v2', http=self.service._http, cache_discovery=False)
        return self.version_service

    def actor(self):
        who = self.identity()
        if not has_full_access(who, self.allowed()):
            raise PermissionError('Sign in with an approved collaborator account.')
        return who['email'].strip().casefold()

    def is_admin(self):
        return self.actor() in {x.strip().casefold() for x in self.admins()}

    def read(self):
        self.actor()
        metadata_request = self.service.files().get(fileId=SOURCE_ID, fields='mimeType,parents,trashed,size', supportsAllDrives=True)
        meta = metadata_request.execute()
        if meta.get('trashed') or FOLDER_ID not in meta.get('parents', []) or meta['mimeType'] != MIME:
            raise PermissionError('Task editing requires the original Excel workbook in the project folder.')
        if int(meta.get('size', 0)) > MAX_BYTES:
            raise ValueError('The workbook exceeds the 25 MB limit.')
        etag = self.versioned().files().get(fileId=SOURCE_ID, fields='etag', supportsAllDrives=True).execute().get('etag')
        request = self.service.files().get_media(fileId=SOURCE_ID, supportsAllDrives=True)
        content = request.execute()
        if len(content) > MAX_BYTES:
            raise ValueError('The workbook exceeds the 25 MB limit.')
        return content, etag

    def list(self):
        actor = self.actor()
        content, _ = self.read()
        tasks, _ = parse_register(content)
        version = sha256(content).hexdigest()
        return [self.adapt(t, version) for t in tasks if self.is_admin() or actor in emails(t['assignee'])]

    @staticmethod
    def adapt(task, version):
        return dict(id=task['row'], title=task['activity'], description=task['output'],
                    assignee=task['assignee'], owner=task['owner'], workstream=task['workstream'],
                    start_date=task['start'].isoformat() if task['start'] else '',
                    due_date=task['end'].isoformat() if task['end'] else '', status=task['status'],
                    progress=100 if task['status'].casefold() in ('completed', 'complete', 'done') and task['progress'] is None else task['progress'], note=task['note'], version=version,
                    reminder_day=task['reminder_day'], reminder_state=task['reminder_state'])

    def write(self, content, etag, change):
        if not etag:
            raise ValueError('Drive did not provide a version tag. Editing is disabled to protect concurrent workbook changes.')
        book = load_workbook(BytesIO(content))
        try:
            sheet = book['Activity Register']
            for row in sheet.iter_rows():
                mapping = {str(cell.value or '').strip().casefold(): cell.column for cell in row if cell.value}
                if {'activity', 'start', 'end', 'owner', 'status'}.issubset(mapping):
                    header_row = row[0].row
                    break
            else:
                raise ValueError('Activity Register headers are missing.')
            def column(name):
                if name.casefold() not in mapping:
                    mapping[name.casefold()] = sheet.max_column + 1
                    sheet.cell(header_row, mapping[name.casefold()], name)
                return mapping[name.casefold()]
            change(sheet, column, header_row)
            out = BytesIO()
            book.save(out)
        finally:
            book.close()
        self.actor()
        request = self.versioned().files().update(fileId=SOURCE_ID,
            media_body=MediaIoBaseUpload(BytesIO(out.getvalue()), mimetype=MIME, resumable=False),
            fields='id', supportsAllDrives=True)
        request.headers['If-Match'] = etag
        try:
            request.execute()
        except Exception as exc:
            if getattr(getattr(exc, 'resp', None), 'status', None) == 412:
                raise ValueError('The workbook changed. Refresh before saving again.') from None
            raise

    def save(self, fields, task_id=None, version=None):
        actor, admin = self.actor(), self.is_admin()
        content, etag = self.read()
        if task_id and sha256(content).hexdigest() != version:
            raise ValueError('The workbook changed. Refresh before saving again.')
        tasks, _ = parse_register(content)
        task = next((t for t in tasks if t['row'] == task_id), None)
        if task_id and (not task or (not admin and actor not in emails(task['assignee']))):
            raise PermissionError('You can only update your assigned tasks.')
        permitted = {'status', 'progress', 'note'}
        if admin:
            permitted |= {'title', 'description', 'assignee', 'owner', 'workstream', 'start_date', 'due_date'}
        if (not task_id and not admin) or set(fields) - permitted:
            raise PermissionError('Only administrators can create tasks or change assignments.')
        current = self.adapt(task, version) if task else dict(status='Not started', progress=0, note='')
        current.update(fields)
        if current['status'] not in STATUSES:
            raise ValueError('Choose a valid status.')
        if current['progress'] is not None and (not isinstance(current['progress'], (int, float)) or not 0 <= current['progress'] <= 100):
            raise ValueError('Progress must be between 0 and 100.')
        if current['status'] == 'Completed':
            current['progress'] = 100
        elif current['progress'] == 100:
            raise ValueError('Choose Completed when progress reaches 100%.')
        if not str(current.get('title', '')).strip() or len(current['title']) > 300:
            raise ValueError('Enter a title of up to 300 characters.')
        if len(current.get('note', '')) > 3000 or len(current.get('description', '')) > 5000:
            raise ValueError('Description or notes exceed their length limit.')
        if not emails(current.get('assignee', '')).issubset({e.strip().casefold() for e in self.allowed()}):
            raise ValueError('Assignee email must contain only approved collaborator addresses.')
        schedule_changed = not task_id or bool({'start_date', 'due_date'} & fields.keys())
        if schedule_changed:
            if parse_date(current['due_date']) < parse_date(current['start_date']):
                raise ValueError('The deadline must be on or after the start date.')
        assignment_changed = admin and (not task_id or 'assignee' in fields) and emails(current.get('assignee','')) != emails(task['assignee'] if task else '')
        if assignment_changed and task and (task['id'] == 'Not recorded' or sum(t['id'] == task['id'] for t in tasks) != 1):
            raise ValueError('Set a unique activity ID in the workbook before changing this assignment.')
        event_ids = []
        def change(sheet, column, header_row):
            target = task_id or max([header_row] + [t['row'] for t in tasks]) + 1
            if admin:
                initialize_tracking(sheet.parent, tasks)
            if not task_id:
                sheet.cell(target, column('ID'), 'T' + datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f'))
            names = dict(title='Activity', description='Key output / milestone', assignee='Assignee email',
                         owner='Owner', workstream='Workstream', start_date='Start', due_date='End',
                         status='Status', progress='Progress (%)', note='Progress notes')
            changed = set(fields)
            if current['status'] == 'Completed':
                changed.add('progress')
            for key in changed:
                cell = sheet.cell(target, column(names[key]))
                value = current[key]
                if key in ('start_date', 'due_date'):
                    cell.value = parse_date(value)
                    cell.number_format = 'dd mmm yyyy'
                else:
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = 's'  # Treat user text literally, including leading '='.
                    if key == 'progress':
                        cell.number_format = '0.##'
            if schedule_changed and 'duration (days)' in [str(c.value or '').casefold() for c in sheet[header_row]]:
                sheet.cell(target, column('Duration (days)'), (parse_date(current['due_date'])-parse_date(current['start_date'])).days+1)
            if assignment_changed:
                activity_id = task['id'] if task else str(sheet.cell(target, column('ID')).value)
                event_ids.extend(record_assignment(sheet.parent, activity_id, current.get('assignee',''), previous_assignee=task['assignee'] if task else ''))
        self.write(content, etag, change)
        return {'notification_events': event_ids}

    def notification_events(self):
        if not self.is_admin():
            raise PermissionError('Administrator access required for notification history.')
        content, _ = self.read()
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            return read_events(book)
        finally:
            book.close()

    def sync_assignments(self):
        """Detect direct workbook edits. First use baselines historical assignments."""
        if not self.is_admin():
            raise PermissionError('Administrator access required for assignment tracking.')
        content, etag = self.read()
        tasks, _ = parse_register(content)
        events = []
        # Skip writes when no assignment changes need tracking.
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            baseline = TRACKING not in book.sheetnames
            previous = {} if baseline else {str(r[0]): emails(str(r[1] or '')) for r in book[TRACKING].iter_rows(min_row=2, max_col=2, values_only=True) if r[0]}
        finally:
            book.close()
        unique = [t for t in tasks if t['id'] != 'Not recorded' and sum(x['id']==t['id'] for x in tasks)==1]
        changed = [t for t in unique if previous.get(t['id'], set()) != emails(t['assignee']) or t['id'] not in previous]
        if not baseline and not changed:
            return []
        def change(sheet, column, header_row):
            if baseline:
                initialize_tracking(sheet.parent, unique)
            else:
                for task in changed:
                    events.extend(record_assignment(sheet.parent, task['id'], task['assignee']))
        self.write(content, etag, change)
        return events

    def claim_assignment(self, event_id):
        if not self.is_admin():
            raise PermissionError('Administrator access required for notification delivery.')
        content, etag = self.read()
        tasks, _ = parse_register(content)
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            event = next((e for e in read_events(book) if e['event_id']==event_id), None)
        finally:
            book.close()
        if not event or event['state'] != 'queued':
            return None
        matches = [t for t in tasks if t['id'] == event['activity_id']]
        task = matches[0] if len(matches)==1 else None
        valid = task and event['recipient'] in emails(task['assignee']) and event['recipient'] in {e.strip().casefold() for e in self.allowed()}
        state = 'claimed' if valid else 'skipped'
        def change(sheet, column, header_row):
            for row in sheet.parent[QUEUE].iter_rows(min_row=2):
                if str(row[0].value)==event_id:
                    row[4].value=state
                    row[5].value=datetime.now(timezone.utc).isoformat()
                    break
        self.write(content, etag, change)
        return {'task': task, 'recipient': event['recipient']} if valid else None

    def finish_assignment(self, event_id, state):
        if not self.is_admin() or state not in ('sent','failed_or_uncertain'):
            raise PermissionError('Administrator access required for notification status updates.')
        content, etag = self.read()
        def change(sheet, column, header_row):
            for row in sheet.parent[QUEUE].iter_rows(min_row=2):
                if str(row[0].value)==event_id and row[4].value=='claimed':
                    row[4].value=state
                    row[5].value=datetime.now(timezone.utc).isoformat()
                    return
            raise ValueError('The notification claim changed. Check delivery status before retrying.')
        self.write(content, etag, change)

    def claim_reminder(self, row, day):
        if not self.is_admin():
            raise PermissionError('Administrator access required for reminder delivery.')
        content, etag = self.read()
        tasks, _ = parse_register(content)
        task = next((t for t in tasks if t['row'] == row), None)
        if not task or task['reminder_day'] == day or task['progress'] == 100:
            return False
        def change(sheet, column, header_row):
            sheet.cell(row, column('Last reminder date'), day)
            sheet.cell(row, column('Last reminder state'), 'claimed')
        self.write(content, etag, change)
        return True
