"""Public team fields and administrator-only edits in the project workbook."""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from uuid import uuid4
import warnings
from openpyxl import load_workbook
from PIL import Image, ImageOps
from googleapiclient.http import MediaIoBaseUpload
from workbook_tasks import WorkbookTasks, MIME
from drive_store import FOLDER_ID, MAX_BYTES
from sheet_dashboard import SOURCE_ID

FIELDS = {'name':'Name','initials':'Initial','affiliation':'Affiliation','expertise':'Experience','position':'Position'}
LIMITS = {'name':160,'initials':20,'affiliation':500,'expertise':3000,'position':300}
PORTRAITS = {'John Humphrey Amuasi':'John_H_Amuasi.png','Anthony Afum-Adjei Awuah':'Anthony_A_A_Awuah.png',
             'Christian Obirikorang':'Christian_Obirikorang.png','Evans Asamoah Adu':'Evans_Asamoah_Adu.png',
             'Afolabi Owoloye':'Afolabi_Owoloye.png','Natalia Shakela':'Natalia_Shakela.png'}


def parse_team(content):
    book=load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        if 'Team_members' not in book.sheetnames:
            raise ValueError('The workbook needs its Team_members tab.')
        members=[];header=None
        for number,row in enumerate(book['Team_members'].iter_rows(values_only=True),1):
            if header is None:
                labels=[str(v or '').strip().casefold() for v in row]
                if {'name','affiliation','position'}.issubset(labels):
                    header={name:i for i,name in enumerate(labels) if name}
                continue
            def get(label):
                index=header.get(label.casefold())
                return str(row[index] or '').strip() if index is not None and index<len(row) else ''
            if not get('Name'):
                continue
            member={key:get(label) for key,label in FIELDS.items()}
            member.update(row=number,profile_id=get('Profile ID') or 'legacy-row-'+str(number),
                          photo_id=get('Photo file ID'),local_portrait=get('Local portrait') or PORTRAITS.get(member['name'],''),featured=get('Featured').casefold() in ('true','yes','1') if get('Featured') else member['name'] in PORTRAITS)
            members.append(member)
        if header is None:
            raise ValueError('Team_members headers were not found.')
        if len({m['profile_id'] for m in members})!=len(members):
            raise ValueError('Team profile IDs must be unique.')
        return members
    finally:
        book.close()


def normalize_portrait(content):
    if not content or len(content)>5*1024*1024:
        raise ValueError('Upload a JPEG or PNG portrait of up to 5 MB.')
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error',Image.DecompressionBombWarning)
            with Image.open(BytesIO(content)) as picture:
                if picture.format not in ('JPEG','PNG') or picture.width*picture.height>16_000_000:
                    raise ValueError('Use a JPEG or PNG image of at most 16 megapixels.')
                picture.load()
                portrait=ImageOps.exif_transpose(picture).convert('RGBA')
                portrait.thumbnail((800,800))
                clean=Image.new('RGB',portrait.size,'white')
                clean.paste(portrait,mask=portrait.getchannel('A'))
                output=BytesIO();clean.save(output,format='PNG')
                return output.getvalue()
    except ValueError:
        raise
    except Exception:
        raise ValueError('The portrait could not be read. Upload a valid JPEG or PNG.') from None


def public_blob(service, file_id, portrait=False):
    # No caller-controlled document endpoint: only the fixed team source and
    # app-uploaded, explicitly published portraits may be read publicly.
    if not portrait and file_id!=SOURCE_ID:
        raise PermissionError('This is not the public team source.')
    meta=service.files().get(fileId=file_id,fields='mimeType,parents,trashed,size,appProperties',supportsAllDrives=True).execute()
    limit=4*1024*1024 if portrait else MAX_BYTES
    if meta.get('trashed') or FOLDER_ID not in meta.get('parents',[]) or int(meta.get('size',0))>limit:
        raise PermissionError('The team resource is unavailable.')
    if portrait:
        if meta.get('mimeType')!='image/png' or meta.get('appProperties',{}).get('purpose')!='public_team_portrait':
            raise PermissionError('This image was not published as a team portrait.')
    elif meta.get('mimeType')!=MIME:
        raise ValueError('The team source must remain an Excel workbook.')
    data=service.files().get_media(fileId=file_id,supportsAllDrives=True).execute()
    if len(data)>limit:
        raise ValueError('The team resource exceeds its size limit.')
    return data


class TeamProfiles(WorkbookTasks):
    def save_profile(self, fields, row=None, version=None, portrait=None):
        if not self.is_admin():
            raise PermissionError('Administrator access is required to edit team profiles.')
        if set(fields)!=(set(FIELDS)|{'featured'}):
            raise ValueError('Use the supported public profile fields.')
        clean={key:str(fields[key]).strip() for key in FIELDS}
        for key,limit in LIMITS.items():
            if len(clean[key])>limit or key in ('name','affiliation','position') and not clean[key]:
                raise ValueError(f'Provide a valid {key} (maximum {limit} characters).')
        if not isinstance(fields['featured'],bool):
            raise ValueError('Choose a valid profile display option.')
        photo=normalize_portrait(portrait) if portrait else None
        content,etag=self.read()
        if not etag:
            raise ValueError('Drive did not provide a version tag. Reload before publishing a profile.')
        if sha256(content).hexdigest()!=version:
            raise ValueError('The workbook changed. Reload profiles before saving; your current text has not been published.')
        members=parse_team(content)
        existing=next((m for m in members if m['row']==row),None)
        if row and not existing:
            raise ValueError('The selected profile no longer exists. Reload the team.')
        photo_id=None
        if photo:
            self.actor()
            photo_id=self.service.files().create(body={'name':'Team_portrait_'+str(uuid4())+'.png','parents':[FOLDER_ID],
                'appProperties':{'purpose':'public_team_portrait'}},media_body=MediaIoBaseUpload(BytesIO(photo),mimetype='image/png'),
                fields='id',supportsAllDrives=True).execute()['id']
        def change(register,column,header_row):
            sheet=register.parent['Team_members']
            for cells in sheet.iter_rows():
                mapping={str(c.value or '').strip().casefold():c.column for c in cells if c.value}
                if {'name','affiliation','position'}.issubset(mapping):
                    team_header=cells[0].row;break
            else:
                raise ValueError('Team headers are missing.')
            def set_value(label,value):
                key=label.casefold()
                if key not in mapping:
                    mapping[key]=sheet.max_column+1
                    sheet.cell(team_header,mapping[key],label)
                cell=sheet.cell(target,mapping[key]);cell.value=value
                if isinstance(value,str):cell.data_type='s'
            target=row or max([team_header]+[m['row'] for m in members])+1
            for key,label in FIELDS.items():set_value(label,clean[key])
            set_value('Profile ID',existing['profile_id'] if existing and not existing['profile_id'].startswith('legacy-row-') else str(uuid4()))
            set_value('Featured','Yes' if fields['featured'] else 'No')
            if existing and existing.get('local_portrait') in PORTRAITS.values():
                set_value('Local portrait',existing['local_portrait'])
            if photo_id:set_value('Photo file ID',photo_id)
            set_value('Updated at',datetime.now(timezone.utc).isoformat())
            set_value('Updated by',self.actor())
        self.write(content,etag,change)
